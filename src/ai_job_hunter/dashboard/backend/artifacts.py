"""
artifacts.py — Base document upload, application queue, and job artifact management.

Handles:
- Parsing uploaded PDF/DOCX/text files into markdown
- Managing the application queue (jobs queued for tailored application)
- Storing and versioning AI-tailored resume and cover letter artifacts
- LLM generation of tailored resumes and cover letters via OpenRouter
- PDF rendering of markdown artifacts
- URL-based artifact lookup for the Chrome extension side panel
"""
from __future__ import annotations

import json
import logging
import os
import re
from functools import lru_cache
from pathlib import Path
from datetime import datetime, timezone
from typing import Any
from urllib.parse import urlparse, urlunparse

import pydantic
import yaml

from ai_job_hunter import settings_service

logger = logging.getLogger(__name__)

_OPENROUTER_BASE_URL = "https://openrouter.ai/api/v1"
_DEFAULT_LLM_MODEL = "z-ai/glm-5.1"
_LLM_PLUGINS = [{"id": "response-healing"}]
_PROMPTS_DIR = Path(__file__).resolve().parents[3].parent / "prompts"


@lru_cache(maxsize=8)
def _load_prompt(filename: str) -> dict:
    path = _PROMPTS_DIR / filename
    if not path.exists():
        raise FileNotFoundError(f"Prompt file not found: {path}")
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    node = next((v for v in data.values() if isinstance(v, dict)), None) or data
    return node


def _render(template: str, variables: dict[str, str]) -> str:
    rendered = template
    for key, value in variables.items():
        rendered = rendered.replace(f"{{{key}}}", value)
    return rendered


class _AtsCritiqueOutput(pydantic.BaseModel):
    """Structured output for ATS resume critique."""
    pass_likelihood: int = pydantic.Field(
        description="Probability 0-100 that an ATS system will pass this resume through to a recruiter."
    )
    missing_keywords: list[str] = pydantic.Field(
        default_factory=list,
        description="Important keywords and skills from the job description that are absent from the resume.",
    )
    weak_sections: list[str] = pydantic.Field(
        default_factory=list,
        description="Names of resume sections that are weak, thin, or missing entirely.",
    )
    suggestions: list[str] = pydantic.Field(
        default_factory=list,
        description="3-5 specific, actionable improvements the candidate should make.",
    )
    revised_resume: str | None = pydantic.Field(
        None,
        description="Improved version of the resume in Markdown with missing keywords woven in naturally. null if no revision is needed.",
    )


# ---------------------------------------------------------------------------
# File parsing
# ---------------------------------------------------------------------------

def parse_uploaded_file(filename: str, content_bytes: bytes, mime_type: str) -> str:
    """
    Parse a PDF, DOCX, or plain text file into a markdown/text string.
    Returns the extracted text content.
    """
    fname_lower = filename.lower()

    if mime_type == "application/pdf" or fname_lower.endswith(".pdf"):
        return _parse_pdf(content_bytes)

    if (mime_type in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                      "application/msword")
            or fname_lower.endswith(".docx") or fname_lower.endswith(".doc")):
        return _parse_docx(content_bytes)

    # Plain text / markdown
    try:
        return content_bytes.decode("utf-8", errors="replace").strip()
    except Exception:
        return ""


def _parse_pdf(content_bytes: bytes) -> str:
    try:
        import io
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content_bytes)) as pdf:
            parts = []
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    parts.append(text.strip())
        return "\n\n".join(parts).strip()
    except ImportError:
        raise RuntimeError("pdfplumber is required for PDF parsing. Run: uv add pdfplumber")
    except Exception as exc:
        raise RuntimeError(f"Failed to parse PDF: {exc}") from exc


def _parse_docx(content_bytes: bytes) -> str:
    try:
        import io
        from docx import Document
        doc = Document(io.BytesIO(content_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs).strip()
    except ImportError:
        raise RuntimeError("python-docx is required for DOCX parsing. Run: uv add python-docx")
    except Exception as exc:
        raise RuntimeError(f"Failed to parse DOCX: {exc}") from exc


# ---------------------------------------------------------------------------
# Base document CRUD
# ---------------------------------------------------------------------------

def list_base_documents(conn: Any) -> list[dict]:
    rows = conn.execute(
        "SELECT id, doc_type, filename, content_md, is_default, created_at FROM base_documents ORDER BY created_at DESC"
    ).fetchall()
    result = []
    for row in rows:
        result.append({
            "id": row[0],
            "doc_type": row[1],
            "filename": row[2],
            "content_md": row[3],
            "is_default": bool(row[4]),
            "created_at": row[5],
        })
    return result


def get_base_document(doc_id: int, conn: Any) -> dict | None:
    row = conn.execute(
        "SELECT id, doc_type, filename, content_md, is_default, created_at FROM base_documents WHERE id = ?",
        (doc_id,)
    ).fetchone()
    if not row:
        return None
    return {
        "id": row[0],
        "doc_type": row[1],
        "filename": row[2],
        "content_md": row[3],
        "is_default": bool(row[4]),
        "created_at": row[5],
    }


def save_base_document(
    doc_type: str,
    filename: str,
    content_md: str,
    content_raw: bytes | None,
    mime_type: str | None,
    conn: Any,
) -> int:
    conn.execute(
        """
        INSERT INTO base_documents (doc_type, filename, content_md, content_raw, mime_type)
        VALUES (?, ?, ?, ?, ?)
        """,
        (doc_type, filename, content_md, content_raw, mime_type),
    )
    conn.commit()
    row = conn.execute("SELECT last_insert_rowid()").fetchone()
    return int(row[0])


def set_default_base_document(doc_id: int, doc_type: str, conn: Any) -> None:
    conn.execute(
        "UPDATE base_documents SET is_default = 0 WHERE doc_type = ?",
        (doc_type,),
    )
    conn.execute(
        "UPDATE base_documents SET is_default = 1 WHERE id = ?",
        (doc_id,),
    )
    conn.commit()


def delete_base_document(doc_id: int, conn: Any) -> None:
    conn.execute("DELETE FROM base_documents WHERE id = ?", (doc_id,))
    conn.commit()


# ---------------------------------------------------------------------------
# Application queue
# ---------------------------------------------------------------------------

def list_queue(conn: Any) -> list[dict]:
    rows = conn.execute(
        """
        SELECT
            q.id, q.job_id, q.status, q.sort_order, q.queued_at, q.processed_at,
            j.title, j.company, j.location, j.ats,
            ms.score
        FROM application_queue q
        JOIN jobs j ON j.id = q.job_id
        LEFT JOIN candidate_profile cp ON cp.id = 1
        LEFT JOIN job_match_scores ms ON ms.job_id = q.job_id AND ms.profile_version = COALESCE(cp.score_version, 1)
        ORDER BY q.sort_order ASC, q.queued_at ASC
        """
    ).fetchall()
    result = []
    for row in rows:
        result.append({
            "id": row[0],
            "job_id": row[1],
            "status": row[2],
            "sort_order": row[3],
            "queued_at": row[4],
            "processed_at": row[5],
            "title": row[6] or "",
            "company": row[7] or "",
            "location": row[8],
            "ats": row[9],
            "match_score": float(row[10]) if row[10] is not None else None,
        })
    return result


def add_to_queue(job_id: str, conn: Any) -> dict:
    existing = conn.execute(
        "SELECT id, job_id, status, sort_order, queued_at FROM application_queue WHERE job_id = ?",
        (job_id,)
    ).fetchone()
    if existing:
        return _queue_row_to_dict(existing, job_id, conn)

    max_row = conn.execute("SELECT COALESCE(MAX(sort_order), -1) FROM application_queue").fetchone()
    next_order = (max_row[0] if max_row else -1) + 1

    conn.execute(
        "INSERT INTO application_queue (job_id, sort_order) VALUES (?, ?)",
        (job_id, next_order),
    )
    conn.commit()

    row = conn.execute(
        "SELECT id, job_id, status, sort_order, queued_at FROM application_queue WHERE job_id = ?",
        (job_id,)
    ).fetchone()
    return _queue_row_to_dict(row, job_id, conn)


def _queue_row_to_dict(row: tuple, job_id: str, conn: Any) -> dict:
    job_row = conn.execute(
        "SELECT title, company, location, ats FROM jobs WHERE id = ?", (job_id,)
    ).fetchone()
    ms_row = conn.execute(
        """
        SELECT ms.score
        FROM job_match_scores ms
        LEFT JOIN candidate_profile cp ON cp.id = 1
        WHERE ms.job_id = ? AND ms.profile_version = COALESCE(cp.score_version, 1)
        """,
        (job_id,),
    ).fetchone()
    return {
        "id": row[0],
        "job_id": row[1],
        "status": row[2],
        "sort_order": row[3],
        "queued_at": row[4],
        "processed_at": None,
        "title": job_row[0] if job_row else "",
        "company": job_row[1] if job_row else "",
        "location": job_row[2] if job_row else None,
        "ats": job_row[3] if job_row else None,
        "match_score": float(ms_row[0]) if ms_row and ms_row[0] is not None else None,
    }


def remove_from_queue(queue_id: int, conn: Any) -> None:
    conn.execute("DELETE FROM application_queue WHERE id = ?", (queue_id,))
    conn.commit()


def update_queue_item(queue_id: int, status: str, conn: Any) -> dict:
    now = datetime.now(timezone.utc).isoformat()
    processed_at = now if status in ("applied", "skipped") else None
    conn.execute(
        "UPDATE application_queue SET status = ?, processed_at = ? WHERE id = ?",
        (status, processed_at, queue_id),
    )
    conn.commit()
    row = conn.execute(
        "SELECT id, job_id, status, sort_order, queued_at FROM application_queue WHERE id = ?",
        (queue_id,)
    ).fetchone()
    if not row:
        return {}
    return _queue_row_to_dict(row, row[1], conn)


def reorder_queue(ids: list[int], conn: Any) -> None:
    for order, qid in enumerate(ids):
        conn.execute(
            "UPDATE application_queue SET sort_order = ? WHERE id = ?",
            (order, qid),
        )
    conn.commit()


# ---------------------------------------------------------------------------
# Job artifacts
# ---------------------------------------------------------------------------

_ARTIFACT_SELECT = """
    SELECT id, job_id, artifact_type, content_md, base_doc_id, version,
           is_active, generated_by, created_at, updated_at, story_ids_used,
           provenance_json, ats_keyword_score, ats_screener_verdict,
           iteration_index, apply_operation_id
    FROM job_artifacts
"""


def get_artifacts_for_job(job_id: str, conn: Any) -> list[dict]:
    rows = conn.execute(
        f"{_ARTIFACT_SELECT} WHERE job_id = ? AND is_active = 1 ORDER BY artifact_type, version DESC",
        (job_id,)
    ).fetchall()
    return [_artifact_row_to_dict(row) for row in rows]


def get_artifact(artifact_id: int, conn: Any) -> dict | None:
    row = conn.execute(
        f"{_ARTIFACT_SELECT} WHERE id = ?",
        (artifact_id,)
    ).fetchone()
    return _artifact_row_to_dict(row) if row else None


def _artifact_row_to_dict(row: tuple) -> dict:
    story_ids: list[int] = []
    if len(row) > 10 and row[10]:
        try:
            story_ids = json.loads(row[10]) or []
        except Exception:
            story_ids = []
    provenance: list[dict] = []
    if len(row) > 11 and row[11]:
        try:
            provenance = json.loads(row[11]) or []
        except Exception:
            provenance = []
    return {
        "id": row[0],
        "job_id": row[1],
        "artifact_type": row[2],
        "content_md": row[3],
        "base_doc_id": row[4],
        "version": row[5],
        "is_active": bool(row[6]),
        "generated_by": row[7],
        "created_at": row[8],
        "updated_at": row[9],
        "story_ids_used": story_ids,
        "provenance": provenance,
        "ats_keyword_score": row[12] if len(row) > 12 else None,
        "ats_screener_verdict": row[13] if len(row) > 13 else None,
        "iteration_index": row[14] if len(row) > 14 else None,
        "apply_operation_id": row[15] if len(row) > 15 else None,
    }


def save_artifact(
    job_id: str,
    artifact_type: str,
    content_md: str,
    base_doc_id: int | None,
    generated_by: str | None,
    conn: Any,
    *,
    story_ids_used: list[int] | None = None,
    provenance: list[dict] | None = None,
    ats_keyword_score: int | None = None,
    ats_screener_verdict: str | None = None,
    iteration_index: int | None = None,
    apply_operation_id: str | None = None,
    make_active: bool = True,
) -> dict:
    if make_active:
        conn.execute(
            "UPDATE job_artifacts SET is_active = 0 WHERE job_id = ? AND artifact_type = ? AND is_active = 1",
            (job_id, artifact_type),
        )

    ver_row = conn.execute(
        "SELECT COALESCE(MAX(version), 0) FROM job_artifacts WHERE job_id = ? AND artifact_type = ?",
        (job_id, artifact_type),
    ).fetchone()
    next_version = (ver_row[0] if ver_row else 0) + 1

    now = datetime.now(timezone.utc).isoformat()
    story_ids_json = json.dumps(story_ids_used or [])
    provenance_json_str = json.dumps(provenance) if provenance is not None else None
    is_active_flag = 1 if make_active else 0
    conn.execute(
        """
        INSERT INTO job_artifacts
            (job_id, artifact_type, content_md, base_doc_id, version, is_active,
             generated_by, created_at, updated_at, story_ids_used,
             provenance_json, ats_keyword_score, ats_screener_verdict,
             iteration_index, apply_operation_id)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (job_id, artifact_type, content_md, base_doc_id, next_version, is_active_flag,
         generated_by, now, now, story_ids_json,
         provenance_json_str, ats_keyword_score, ats_screener_verdict,
         iteration_index, apply_operation_id),
    )
    conn.commit()

    row = conn.execute(
        f"{_ARTIFACT_SELECT} WHERE job_id = ? AND artifact_type = ? AND version = ?",
        (job_id, artifact_type, next_version),
    ).fetchone()
    return _artifact_row_to_dict(row)


def update_artifact(artifact_id: int, content_md: str, conn: Any) -> dict:
    now = datetime.now(timezone.utc).isoformat()
    conn.execute(
        "UPDATE job_artifacts SET content_md = ?, updated_at = ? WHERE id = ?",
        (content_md, now, artifact_id),
    )
    conn.commit()
    row = conn.execute(f"{_ARTIFACT_SELECT} WHERE id = ?", (artifact_id,)).fetchone()
    return _artifact_row_to_dict(row) if row else {}


# ---------------------------------------------------------------------------
# LLM generation
# ---------------------------------------------------------------------------

def _load_job_context(job_id: str, conn: Any) -> dict:
    row = conn.execute(
        """
        SELECT j.title, j.company, j.location, j.ats,
               e.work_mode, e.seniority, e.role_family,
               e.years_exp_min, e.years_exp_max,
               e.required_skills, e.preferred_skills,
               e.formatted_description,
               j.description
        FROM jobs j
        LEFT JOIN job_enrichments e ON e.job_id = j.id
        WHERE j.id = ?
        """,
        (job_id,)
    ).fetchone()
    if not row:
        raise ValueError(f"Job {job_id} not found")

    def parse_json_list(val: Any) -> list:
        if not val:
            return []
        try:
            return json.loads(val)
        except Exception:
            return []

    return {
        "title": row[0] or "",
        "company": row[1] or "",
        "location": row[2] or "",
        "ats": row[3] or "",
        "work_mode": row[4] or "",
        "seniority": row[5] or "",
        "role_family": row[6] or "",
        "years_exp_min": row[7],
        "years_exp_max": row[8],
        "required_skills": parse_json_list(row[9]),
        "preferred_skills": parse_json_list(row[10]),
        "description": (row[11] or row[12] or "")[:6000],
    }


def _load_profile_context(conn: Any) -> dict:
    row = conn.execute(
        """
        SELECT years_experience, skills, desired_job_titles,
               full_name, email, phone, linkedin_url, portfolio_url, city, country,
               narrative_intent
        FROM candidate_profile WHERE id = 1
        """
    ).fetchone()
    if not row:
        return {}

    def parse_json_list(val: Any) -> list:
        if not val:
            return []
        try:
            return json.loads(val)
        except Exception:
            return []

    return {
        "years_experience": row[0] or 0,
        "skills": parse_json_list(row[1]),
        "desired_job_titles": parse_json_list(row[2]),
        "full_name": row[3] or "",
        "email": row[4] or "",
        "phone": row[5] or "",
        "linkedin_url": row[6] or "",
        "portfolio_url": row[7] or "",
        "city": row[8] or "",
        "country": row[9] or "",
        "narrative_intent": row[10] or "",
    }


def load_story_context_for_generation(job_id: str, conn: Any) -> tuple[str, list[int]]:
    """
    Return (story_block_text, story_ids) for grounding artifact generation.

    Tries semantic similarity first (via embeddings). Falls back to top accepted
    stories by importance when embeddings are unavailable.
    """
    # Try semantic path
    try:
        from ai_job_hunter.dashboard.backend.embeddings import get_relevant_stories_for_job
        stories = get_relevant_stories_for_job(job_id, conn, top_k=5)
    except Exception:
        stories = []

    # Fallback: top accepted stories by importance
    if not stories:
        rows = conn.execute(
            """
            SELECT id, title, kind, narrative, role_context, skills, outcomes, importance
            FROM user_stories
            WHERE draft = 0
            ORDER BY importance DESC, created_at DESC
            LIMIT 5
            """
        ).fetchall()
        for r in rows:
            skills_list: list[str] = []
            outcomes_list: list[str] = []
            try:
                skills_list = json.loads(r[5] or "[]") or []
            except Exception:
                pass
            try:
                outcomes_list = json.loads(r[6] or "[]") or []
            except Exception:
                pass
            stories.append({
                "id": r[0],
                "title": r[1],
                "kind": r[2],
                "narrative": r[3],
                "role_context": r[4],
                "skills": skills_list,
                "outcomes": outcomes_list,
                "importance": r[7],
            })

    if not stories:
        return "", []

    story_ids = [int(s["id"]) for s in stories]
    lines = [
        "STORY BANK — YOUR VERIFIED REAL EXPERIENCES:",
        "Every achievement, metric, and claim in the generated document MUST come",
        "from these stories or the Base Resume. Do NOT invent or embellish.",
        "",
    ]
    for i, s in enumerate(stories, 1):
        lines.append(f"--- [{i}] {s['title']} ({s['kind']}) ---")
        if s.get("role_context"):
            lines.append(f"Context: {s['role_context']}")
        if s.get("narrative"):
            lines.append(f"Your experience: {str(s['narrative'])[:800]}")
        if s.get("skills"):
            lines.append("Skills demonstrated: " + ", ".join(str(x) for x in s["skills"][:12]))
        if s.get("outcomes"):
            lines.append("Key outcomes: " + " | ".join(str(x) for x in s["outcomes"][:5]))
        lines.append("")

    return "\n".join(lines), story_ids


def _build_llm() -> Any:
    api_key = settings_service.get("OPENROUTER_API_KEY").strip()
    if not api_key:
        raise RuntimeError("OPENROUTER_API_KEY not set")
    model = settings_service.get("LLM_MODEL").strip()
    try:
        from langchain_openai import ChatOpenAI
    except ImportError:
        raise RuntimeError("langchain-openai is required. Run: uv add langchain-openai")
    return ChatOpenAI(
        model=model,
        api_key=api_key,
        base_url=_OPENROUTER_BASE_URL,
        temperature=0.3,
        max_tokens=3000,
        extra_body={"plugins": _LLM_PLUGINS},
    )


def _call_llm(prompt: str, system: str) -> str:
    try:
        from langchain_core.messages import HumanMessage, SystemMessage
    except ImportError:
        raise RuntimeError("langchain-openai is required. Run: uv add langchain-openai")
    llm = _build_llm()
    response = llm.invoke([SystemMessage(content=system), HumanMessage(content=prompt)])
    return str(getattr(response, "content", "") or "").strip()


async def _call_llm_astream(prompt: str, system: str):
    """Async generator yielding str token chunks from the LLM."""
    try:
        from langchain_core.messages import HumanMessage, SystemMessage
    except ImportError:
        raise RuntimeError("langchain-openai is required. Run: uv add langchain-openai")
    llm = _build_llm()
    async for chunk in llm.astream([SystemMessage(content=system), HumanMessage(content=prompt)]):
        token = str(getattr(chunk, "content", "") or "")
        if token:
            yield token


def generate_tailored_resume(
    job_id: str,
    base_doc_id: int,
    conn: Any,
    *,
    story_context: str = "",
    story_ids_used: list[int] | None = None,
) -> tuple[str, list[int]]:
    """Returns (content_md, story_ids_used)."""
    job = _load_job_context(job_id, conn)
    profile = _load_profile_context(conn)
    base_doc = get_base_document(base_doc_id, conn)
    if not base_doc:
        raise ValueError(f"Base document {base_doc_id} not found")

    # Load story context if not pre-loaded
    _story_ids = story_ids_used or []
    if not story_context:
        story_context, _story_ids = load_story_context_for_generation(job_id, conn)

    req_skills = ", ".join(job["required_skills"][:10]) if job["required_skills"] else "Not specified"
    pref_skills = ", ".join(job["preferred_skills"][:10]) if job["preferred_skills"] else "Not specified"
    profile_skills = ", ".join(profile.get("skills", [])[:20]) if profile.get("skills") else "Not specified"
    seniority = job["seniority"] or "Not specified"
    years_range = ""
    if job["years_exp_min"] is not None:
        years_range = f"{job['years_exp_min']}"
        if job["years_exp_max"] is not None:
            years_range += f"-{job['years_exp_max']}"
        years_range += " years"

    tpl = _load_prompt("resume_tailoring.yaml")
    story_section = f"\n{story_context}\n" if story_context else ""
    grounding_rules = ("\n" + tpl["grounding_rules"].strip() + "\n") if story_context else ""
    vars_ = {
        "job_title": str(job.get("title") or ""),
        "job_company": str(job.get("company") or ""),
        "job_location": str(job.get("location") or ""),
        "job_seniority": seniority,
        "job_work_mode": str(job.get("work_mode") or "Not specified"),
        "job_required_skills": req_skills,
        "job_preferred_skills": pref_skills,
        "job_years_range": years_range or "Not specified",
        "job_description": str(job.get("description") or "")[:2500],
        "story_section": story_section,
        "candidate_name": str(profile.get("full_name") or ""),
        "candidate_email": str(profile.get("email") or ""),
        "candidate_phone": str(profile.get("phone") or ""),
        "candidate_linkedin": str(profile.get("linkedin_url") or ""),
        "candidate_location": f"{profile.get('city', '')}, {profile.get('country', '')}",
        "candidate_years_exp": str(profile.get("years_experience") or 0),
        "candidate_skills": profile_skills,
        "base_resume": str(base_doc.get("content_md") or "")[:3000],
        "grounding_rules": grounding_rules,
    }
    system = tpl["system"].strip()
    prompt = _render(tpl["user"], vars_)
    content_md = _call_llm(prompt, system)
    return content_md, _story_ids


def generate_cover_letter(
    job_id: str,
    base_doc_id: int,
    conn: Any,
    *,
    story_context: str = "",
    story_ids_used: list[int] | None = None,
) -> tuple[str, list[int]]:
    """Returns (content_md, story_ids_used)."""
    job = _load_job_context(job_id, conn)
    profile = _load_profile_context(conn)
    base_doc = get_base_document(base_doc_id, conn)
    if not base_doc:
        raise ValueError(f"Base document {base_doc_id} not found")

    _story_ids = story_ids_used or []
    if not story_context:
        story_context, _story_ids = load_story_context_for_generation(job_id, conn)

    req_skills = ", ".join(job["required_skills"][:8]) if job["required_skills"] else "Not specified"
    profile_skills = ", ".join(profile.get("skills", [])[:15]) if profile.get("skills") else "Not specified"

    tpl = _load_prompt("cover_letter.yaml")
    story_section = f"\n{story_context}\n" if story_context else ""
    grounding_rules = ("\n" + tpl["grounding_rules"].strip() + "\n") if story_context else ""
    vars_ = {
        "job_title": str(job.get("title") or ""),
        "job_company": str(job.get("company") or ""),
        "job_location": str(job.get("location") or ""),
        "job_required_skills": req_skills,
        "job_description": str(job.get("description") or "")[:1800],
        "story_section": story_section,
        "candidate_name": str(profile.get("full_name") or ""),
        "candidate_email": str(profile.get("email") or ""),
        "candidate_phone": str(profile.get("phone") or ""),
        "candidate_linkedin": str(profile.get("linkedin_url") or ""),
        "candidate_location": f"{profile.get('city', '')}, {profile.get('country', '')}",
        "candidate_years_exp": str(profile.get("years_experience") or 0),
        "candidate_skills": profile_skills,
        "base_resume": str(base_doc.get("content_md") or "")[:1500],
        "grounding_rules": grounding_rules,
    }
    system = tpl["system"].strip()
    prompt = _render(tpl["user"], vars_)
    content_md = _call_llm(prompt, system)
    return content_md, _story_ids


async def generate_tailored_resume_astream(
    job_id: str,
    base_doc_id: int,
    conn: Any,
    *,
    story_context: str = "",
):
    """Async generator yielding resume token chunks."""
    job = _load_job_context(job_id, conn)
    profile = _load_profile_context(conn)
    base_doc = get_base_document(base_doc_id, conn)
    if not base_doc:
        raise ValueError(f"Base document {base_doc_id} not found")

    req_skills = ", ".join(job["required_skills"][:10]) if job["required_skills"] else "Not specified"
    pref_skills = ", ".join(job["preferred_skills"][:10]) if job["preferred_skills"] else "Not specified"
    profile_skills = ", ".join(profile.get("skills", [])[:20]) if profile.get("skills") else "Not specified"
    seniority = job["seniority"] or "Not specified"
    years_range = ""
    if job["years_exp_min"] is not None:
        years_range = f"{job['years_exp_min']}"
        if job["years_exp_max"] is not None:
            years_range += f"-{job['years_exp_max']}"
        years_range += " years"

    tpl = _load_prompt("resume_tailoring.yaml")
    story_section = f"\n{story_context}\n" if story_context else ""
    grounding_rules = ("\n" + tpl["grounding_rules"].strip() + "\n") if story_context else ""
    vars_ = {
        "job_title": str(job.get("title") or ""),
        "job_company": str(job.get("company") or ""),
        "job_location": str(job.get("location") or ""),
        "job_seniority": seniority,
        "job_work_mode": str(job.get("work_mode") or "Not specified"),
        "job_required_skills": req_skills,
        "job_preferred_skills": pref_skills,
        "job_years_range": years_range or "Not specified",
        "job_description": str(job.get("description") or "")[:2500],
        "story_section": story_section,
        "candidate_name": str(profile.get("full_name") or ""),
        "candidate_email": str(profile.get("email") or ""),
        "candidate_phone": str(profile.get("phone") or ""),
        "candidate_linkedin": str(profile.get("linkedin_url") or ""),
        "candidate_location": f"{profile.get('city', '')}, {profile.get('country', '')}",
        "candidate_years_exp": str(profile.get("years_experience") or 0),
        "candidate_skills": profile_skills,
        "base_resume": str(base_doc.get("content_md") or "")[:3000],
        "grounding_rules": grounding_rules,
    }
    system = tpl["system"].strip()
    prompt = _render(tpl["user"], vars_)
    async for token in _call_llm_astream(prompt, system):
        yield token


async def generate_cover_letter_astream(
    job_id: str,
    base_doc_id: int,
    conn: Any,
    *,
    story_context: str = "",
):
    """Async generator yielding cover letter token chunks."""
    job = _load_job_context(job_id, conn)
    profile = _load_profile_context(conn)
    base_doc = get_base_document(base_doc_id, conn)
    if not base_doc:
        raise ValueError(f"Base document {base_doc_id} not found")

    req_skills = ", ".join(job["required_skills"][:8]) if job["required_skills"] else "Not specified"
    profile_skills = ", ".join(profile.get("skills", [])[:15]) if profile.get("skills") else "Not specified"

    tpl = _load_prompt("cover_letter.yaml")
    story_section = f"\n{story_context}\n" if story_context else ""
    grounding_rules = ("\n" + tpl["grounding_rules"].strip() + "\n") if story_context else ""
    vars_ = {
        "job_title": str(job.get("title") or ""),
        "job_company": str(job.get("company") or ""),
        "job_location": str(job.get("location") or ""),
        "job_required_skills": req_skills,
        "job_description": str(job.get("description") or "")[:1800],
        "story_section": story_section,
        "candidate_name": str(profile.get("full_name") or ""),
        "candidate_email": str(profile.get("email") or ""),
        "candidate_phone": str(profile.get("phone") or ""),
        "candidate_linkedin": str(profile.get("linkedin_url") or ""),
        "candidate_location": f"{profile.get('city', '')}, {profile.get('country', '')}",
        "candidate_years_exp": str(profile.get("years_experience") or 0),
        "candidate_skills": profile_skills,
        "base_resume": str(base_doc.get("content_md") or "")[:1500],
        "grounding_rules": grounding_rules,
    }
    system = tpl["system"].strip()
    prompt = _render(tpl["user"], vars_)
    async for token in _call_llm_astream(prompt, system):
        yield token


# ---------------------------------------------------------------------------
# ATS critique
# ---------------------------------------------------------------------------

def critique_resume_for_ats(job_id: str, resume_md: str, conn: Any) -> dict:
    """
    Critique a resume's ATS compatibility against the given job.

    Returns a dict with:
      - pass_likelihood: int 0-100
      - missing_keywords: list[str]
      - weak_sections: list[str]
      - suggestions: list[str]
      - revised_resume: str | None
    """
    try:
        from langchain_core.messages import HumanMessage, SystemMessage
    except ImportError:
        raise RuntimeError("langchain-openai is required. Run: uv add langchain-openai")

    job = _load_job_context(job_id, conn)
    req_skills = ", ".join(job["required_skills"][:15]) if job["required_skills"] else "Not specified"
    pref_skills = ", ".join(job["preferred_skills"][:10]) if job["preferred_skills"] else "Not specified"

    tpl = _load_prompt("ats_critique.yaml")
    vars_ = {
        "job_title": str(job.get("title") or ""),
        "job_company": str(job.get("company") or ""),
        "job_required_skills": req_skills,
        "job_preferred_skills": pref_skills,
        "job_description": str(job.get("description") or "")[:2000],
        "resume_md": resume_md[:3000],
    }
    system = tpl["system"].strip()
    prompt = _render(tpl["user"], vars_)

    llm = _build_llm()
    structured = llm.with_structured_output(_AtsCritiqueOutput, method="json_schema", strict=True)
    try:
        result = structured.invoke([SystemMessage(content=system), HumanMessage(content=prompt)])
    except Exception:
        logger.warning("ATS critique LLM call failed", exc_info=True)
        return {
            "pass_likelihood": 50,
            "missing_keywords": [],
            "weak_sections": [],
            "suggestions": ["Unable to generate critique — please try again."],
            "revised_resume": None,
        }

    return {
        "pass_likelihood": max(0, min(100, int(result.pass_likelihood or 50))),
        "missing_keywords": [str(k) for k in (result.missing_keywords or [])[:20]],
        "weak_sections": [str(s) for s in (result.weak_sections or [])[:10]],
        "suggestions": [str(s) for s in (result.suggestions or [])[:10]],
        "revised_resume": str(result.revised_resume) if result.revised_resume else None,
    }


# ---------------------------------------------------------------------------
# PDF rendering
# ---------------------------------------------------------------------------

def render_artifact_pdf(content_md: str) -> bytes:
    try:
        from fpdf import FPDF
        from fpdf.enums import XPos, YPos
    except ImportError:
        raise RuntimeError("fpdf2 is required. Run: uv add fpdf2")

    # Convert markdown to plain-text-like lines for fpdf
    # We parse the markdown structure manually for reliable rendering
    lines = content_md.split("\n")

    pdf = FPDF()
    pdf.add_page()
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(auto=True, margin=20)

    def add_text_line(text: str, size: int = 10, bold: bool = False, italic: bool = False, ln: bool = True) -> None:
        style = ""
        if bold:
            style += "B"
        if italic:
            style += "I"
        pdf.set_font("Helvetica", style=style, size=size)
        safe_text = text.encode("latin-1", errors="replace").decode("latin-1")
        if ln:
            pdf.cell(0, size * 0.5 + 2, safe_text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        else:
            pdf.cell(0, size * 0.5 + 2, safe_text)

    for line in lines:
        stripped = line.strip()

        if stripped.startswith("### "):
            add_text_line(stripped[4:], size=11, bold=True)
        elif stripped.startswith("## "):
            add_text_line(stripped[3:], size=13, bold=True)
        elif stripped.startswith("# "):
            add_text_line(stripped[2:], size=15, bold=True)
        elif stripped.startswith("---") or stripped.startswith("***"):
            pdf.set_draw_color(200, 200, 200)
            pdf.line(pdf.get_x(), pdf.get_y(), pdf.get_x() + 170, pdf.get_y())
            pdf.ln(3)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            bullet_text = stripped[2:]
            # Strip inline bold/italic markers
            bullet_text = re.sub(r"\*\*(.*?)\*\*", r"\1", bullet_text)
            bullet_text = re.sub(r"\*(.*?)\*", r"\1", bullet_text)
            pdf.set_font("Helvetica", size=10)
            safe_text = ("  \u2022 " + bullet_text).encode("latin-1", errors="replace").decode("latin-1")
            pdf.cell(0, 7, safe_text, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        elif stripped == "":
            pdf.ln(3)
        else:
            # Inline bold/italic stripping for regular text
            clean = re.sub(r"\*\*(.*?)\*\*", r"\1", stripped)
            clean = re.sub(r"\*(.*?)\*", r"\1", clean)
            add_text_line(clean, size=10)

    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# URL-based artifact lookup (for Chrome extension side panel)
# ---------------------------------------------------------------------------

_APPLY_SUFFIXES = re.compile(r"(/apply(?:ication)?/?|/form/?|/#apply/?|/\?.*)?$", re.IGNORECASE)


def _normalize_job_url(url: str) -> str:
    """Strip application-page suffixes to get the canonical job listing URL."""
    # Remove query string and fragment
    try:
        parsed = urlparse(url)
        clean = urlunparse(parsed._replace(query="", fragment=""))
    except Exception:
        clean = url
    # Strip trailing /apply, /application, /form etc
    clean = _APPLY_SUFFIXES.sub("", clean).rstrip("/")
    return clean


def find_job_by_application_url(url: str, conn: Any) -> str | None:
    """
    Find a job_id by matching the application page URL against stored job URLs.
    Tries: exact match → normalized match → prefix match.
    Returns job_id or None.
    """
    normalized = _normalize_job_url(url)

    # Exact match
    row = conn.execute("SELECT id FROM jobs WHERE url = ?", (url,)).fetchone()
    if row:
        return row[0]

    # Normalized match
    row = conn.execute("SELECT id FROM jobs WHERE url = ?", (normalized,)).fetchone()
    if row:
        return row[0]

    # Prefix match: stored URL is a prefix of the incoming URL
    # This handles: stored = .../jobs/12345, incoming = .../jobs/12345/apply
    row = conn.execute(
        "SELECT id FROM jobs WHERE ? LIKE url || '%' ORDER BY LENGTH(url) DESC LIMIT 1",
        (normalized,)
    ).fetchone()
    if row:
        return row[0]

    return None


def get_artifacts_by_url(url: str, conn: Any) -> dict:
    """
    Return active artifacts (resume + cover_letter) for the job matching url.
    Used by the Chrome extension side panel GET /api/artifacts/by-url endpoint.
    """
    job_id = find_job_by_application_url(url, conn)
    if not job_id:
        return {"job_info": None, "resume": None, "cover_letter": None}

    job_row = conn.execute(
        "SELECT title, company, location FROM jobs WHERE id = ?", (job_id,)
    ).fetchone()
    job_info = {
        "title": job_row[0] if job_row else "",
        "company": job_row[1] if job_row else "",
        "location": job_row[2] if job_row else None,
        "job_id": job_id,
    } if job_row else None

    artifacts = get_artifacts_for_job(job_id, conn)
    resume = next((a for a in artifacts if a["artifact_type"] == "resume"), None)
    cover_letter = next((a for a in artifacts if a["artifact_type"] == "cover_letter"), None)

    return {
        "job_info": job_info,
        "resume": resume,
        "cover_letter": cover_letter,
    }
